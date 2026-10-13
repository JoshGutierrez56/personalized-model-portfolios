"""Create aggregate-only investment oversight workbook, memo and interview guide."""
from __future__ import annotations

from datetime import datetime, timezone
import hashlib
import html
import json
from pathlib import Path
import re
from urllib.parse import quote

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_LEFT, TA_RIGHT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (SimpleDocTemplate, Paragraph, Spacer, Table,
                                TableStyle, PageBreak, Preformatted, KeepTogether)
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

from control_metadata import describe_control
from release_evidence import validate_release_bundle

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "outputs/oversight"
NAVY = "142D42"
TEAL = "006B65"
GRAY = "EEF2F5"
INK = "182D3E"


def read(name):
    return json.loads((OUT / name).read_text(encoding="utf-8"))


def flat(value):
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return value


def tab(wb, name, rows, title, subtitle):
    ws = wb.create_sheet(name)
    ws.append([title]); ws.append([subtitle]); ws.append([])
    rows = list(rows)
    if not rows:
        ws.append(["No rows"])
    else:
        keys = list(rows[0])
        ws.append(keys)
        for row in rows:
            ws.append([flat(row.get(k)) for k in keys])
        ws.auto_filter.ref = f"A4:{get_column_letter(len(keys))}{ws.max_row}"
    ws.freeze_panes = "B5"
    for cell in ws[1]:
        cell.font = Font(name="Calibri", size=18, bold=True, color=NAVY)
    for cell in ws[2]:
        cell.font = Font(name="Calibri", size=10, italic=True, color="526574")
    for cell in ws[4]:
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    for row in ws.iter_rows(min_row=5):
        for cell in row:
            cell.font = Font(name="Calibri", size=11, color=INK)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if cell.row % 2 == 1:
                cell.fill = PatternFill("solid", fgColor="F3F6F8")
            if isinstance(cell.value, float):
                cell.number_format = "0.0000"
    for j in range(1, ws.max_column + 1):
        lengths = [len(str(ws.cell(i, j).value or "")) for i in range(4, min(ws.max_row, 100) + 1)]
        ws.column_dimensions[get_column_letter(j)].width = min(55, max(16, max(lengths, default=18) + 2))
        if ws.cell(4, j).value in {"max_abs_error", "tolerance", "active_linking_residual"}:
            for row_number in range(5, ws.max_row + 1):
                ws.cell(row_number, j).number_format = "0.00E+00"
    ws.row_dimensions[1].height = 30
    ws.row_dimensions[4].height = 32
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:4"
    ws.oddFooter.center.text = "Portfolio Oversight | Aggregate research / not live orders"
    return ws


def workbook():
    bundle = validate_release_bundle(ROOT)
    baseline, incidents, receipt = bundle["baseline"], bundle["incidents"], bundle["receipt"]
    scenarios = pd.DataFrame(bundle["scenarios"])
    wb = Workbook(); wb.remove(wb.active)
    wb.properties.creator = "Josh Gutierrez"
    wb.properties.title = "Portfolio Validation & Investment Oversight"
    wb.properties.subject = "Independent aggregate research validation and deliberate fault-injection evidence"
    overview = [{"item": "Evidence release decision", "value": receipt["status"], "interpretation": "Technical validation only; not investment or live-trade approval."},
        {"item": "Baseline", "value": baseline["status"], "interpretation": "Separately implemented challenger using the saved licensed research."},
        {"item": "Injected faults detected", "value": f'{receipt["detected_cases"]}/{receipt["injected_cases"]}', "interpretation": "Controlled tests, not real operational incidents."},
        {"item": "Corrected original-input retest", "value": receipt["corrected_retest_status"], "interpretation": "Original source hashes remain unchanged."},
        {"item": "Investment conclusion", "value": "Do not spend tracking-risk budget solely to save tiny ETF costs", "interpretation": "For No gold, 1.10bp/year modeled savings accompanies about 27bp more realized TE."},
        {"item": "Reproducibility", "value": receipt["run_id"], "interpretation": "See RunEvidence; full empirical rerun requires own licensed inputs."}]
    tab(wb, "Decision", overview, "Portfolio oversight | decision brief", "Actual WRDS-backed aggregate research. Hypothetical accounts. Already-exposed 2014–2025 history.")
    controls = [{**check, **describe_control(check["id"])} for check in baseline.get("checks", [])]
    tab(wb, "Controls", controls, "Independent control results", "Residual and tolerance share the explicit unit column. PASS is technical consistency, not investment approval.")
    cases = [{"id": c["id"], "case": c["title"], "classification": c["classification"], "observed": c["observed_status"],
        "retest": c["corrected_retest"], "closure": c["closure"], "root_cause": c["root_cause"],
        "corrective_action": c["corrective_action"], "suggested_owner": c["suggested_operational_owner"],
        "acceptance_criterion": c["acceptance_criterion"]} for c in incidents["cases"]]
    tab(wb, "Exceptions", cases, "Exception register | controlled experiments", "Suggested functional owners are workflow examples, not a claim of a managed production team.")
    cols = ["client", "method", "gross_te_pct", "annual_one_way_turnover_pct", "annual_incurred_cost_bps", "net_cagr_pct", "terminal_nav"]
    ws = tab(wb, "RiskCost", scenarios[cols].to_dict("records"), "Risk and implementation economics", "TE/turnover/CAGR in %. Cost in annualized basis points of account value. Terminal NAV in USD.")
    summary = scenarios[scenarios.client.eq("No gold")].set_index("method")
    a, b = summary.loc["Minimum risk"], summary.loc["Lower turnover"]
    difference = float(a.annual_incurred_cost_bps - b.annual_incurred_cost_bps)
    tab(wb, "NoGoldDecision", [
        {"metric": "Annual one-way turnover (%)", "minimum_risk": a.annual_one_way_turnover_pct, "lower_turnover": b.annual_one_way_turnover_pct},
        {"metric": "Realized gross TE (%)", "minimum_risk": a.gross_te_pct, "lower_turnover": b.gross_te_pct},
        {"metric": "Annualized cost (bp NAV)", "minimum_risk": a.annual_incurred_cost_bps, "lower_turnover": b.annual_incurred_cost_bps},
        {"metric": "Annual cost per constant $1m (USD)", "minimum_risk": a.annual_incurred_cost_bps * 100, "lower_turnover": b.annual_incurred_cost_bps * 100}],
        "Would I authorize the lower-turnover choice?", "Conditional choice, not automatic approval: inspect mandate, cost assumptions and the price of tracking risk.")
    tab(wb, "Stress", bundle["stress"], "Historical episode replay", "Preselected familiar periods within already-exposed history. Returns exclude start-close date, include end-close date.")
    agg = baseline.get("aggregate", {})
    tab(wb, "Reconciliation", [{"metric": k, "value": flat(v)} for k, v in agg.items() if k not in {"attribution_totals", "path_metrics"}],
        "Reconciliation | aggregate residuals", "Detailed security-day ledgers are intentionally excluded from this public workbook.")
    attr = agg.get("attribution_totals", [])
    if agg.get("path_metrics"):
        tab(wb, "PathMetrics", agg["path_metrics"], "Reconstructed path metrics", "One aggregate row per account/method. *_pct fields are percentages; *_bps are basis points. Tracking errors and bound drift are fractions; 0.01 = 1%. NAV/cost totals are USD.")
    if isinstance(attr, list) and attr and isinstance(attr[0], dict):
        totals = [{k: v for k, v in row.items() if not isinstance(v, dict)} for row in attr]
        tab(wb, "LinkedTotals", totals, "Cumulative contribution reconciliation", "Initial-capital-scaled dollar P&L, in decimals. 0.01 = 1 percentage point of cumulative return.")
        detail = [{"client": row["client"], "method": row["method"], "asset": ticker,
                   "asset_pnl_over_initial_capital": value,
                   "benchmark_pnl_over_initial_capital": row["linked_benchmark_pnl"][ticker],
                   "active_asset_pnl_over_initial_capital": row["linked_active_asset_pnl"][ticker]}
                  for row in attr for ticker, value in row["linked_asset_pnl"].items()]
        tab(wb, "Contributions", detail, "Linked asset contributions | full-history aggregates", "Not Brinson/factor attribution or a sum of daily active returns. Add linked cost to summed active contributions to reconcile terminal wealth difference / initial capital.")
    tab(wb, "RunEvidence", [{"key": k, "value": flat(v)} for k, v in receipt.items()],
        "Run provenance and evidence", "Relative paths/hashes identify local research inputs without distributing those inputs.")
    fx = wb.create_sheet("WorkedExample")
    fx.append(["Hand-worked trade accounting | artificial teaching fixture"])
    fx.append(["These numbers are fabricated for checking formulas, NOT empirical research or an actual account."])
    fixture = [("Initial A", 600), ("Initial B", 400), ("A return", .10), ("B return", -.05),
               ("Cost rate", .001), ("Beginning NAV", "=B4+B5"), ("Pretrade A", "=B4*(1+B6)"),
               ("Pretrade B", "=B5*(1+B7)"), ("Pretrade NAV", "=B10+B11"),
               ("Fee (one sell, one buy; 50/50 target)", "=B8*(B10-B11)"),
               ("Posttrade NAV", "=B12-B13"), ("Posttrade A", "=B14/2"), ("Posttrade B", "=B14/2"),
               ("Trade A", "=B15-B10"), ("Trade B", "=B16-B11"),
               ("Funding residual", "=B17+B18+B13"), ("Gross return", "=B12/B9-1"),
               ("Net return", "=B14/B9-1"), ("Cost contribution", "=-B13/B9"),
               ("Return residual", "=B20+B22-B21")]
    fx.append([])
    for row in fixture: fx.append(row)
    fx.column_dimensions["A"].width = 65; fx.column_dimensions["B"].width = 24
    for row in fx.iter_rows(min_row=4):
        row[0].font = Font(name="Calibri", size=11, color=INK)
        row[1].number_format = "0.000000"
        row[1].fill = PatternFill("solid", fgColor="E9F4F2" if isinstance(row[1].value, str) else "EEF2F5")
    fx.freeze_panes = "B4"; fx.sheet_view.showGridLines = False
    from openpyxl.workbook.properties import CalcProperties
    wb.calculation = CalcProperties(calcId=191029, fullCalcOnLoad=True)
    path = OUT / "Portfolio_Oversight.xlsx"
    wb.save(path)
    check = load_workbook(path, read_only=False, data_only=False)
    assert len(check.worksheets) >= 9 and check["WorkedExample"]["B23"].value == "=B20+B22-B21"
    return {"path": path.name, "sheets": check.sheetnames, "formula_note": "Open in Excel/LibreOffice to recalculate teaching formulas; research figures are materialized values."}


def _font_files():
    """Choose complete Unicode families, never silently fall back to base-14 PDF fonts.

    Matplotlib is an existing pinned dependency and distributes free DejaVu fonts.
    Those are the portable default; no network/font installation occurs here.
    """
    try:
        from matplotlib import get_data_path
        dejavu = Path(get_data_path()) / "fonts/ttf"
        portable_files = [dejavu / name for name in
                          ("DejaVuSans.ttf", "DejaVuSans-Bold.ttf", "DejaVuSans-Oblique.ttf",
                           "DejaVuSans-BoldOblique.ttf", "DejaVuSansMono.ttf")]
        if all(path.is_file() for path in portable_files):
            return portable_files, "DejaVu Sans", "DejaVu Sans Mono"
    except ImportError:
        pass
    windows = Path("C:/Windows/Fonts")
    windows_files = [windows / name for name in
                     ("arial.ttf", "arialbd.ttf", "ariali.ttf", "arialbi.ttf", "consola.ttf")]
    if all(path.is_file() for path in windows_files):
        return windows_files, "Arial", "Consolas"
    raise RuntimeError("Unicode report fonts unavailable. Install the pinned matplotlib dependency "
                       "with its bundled DejaVu Sans and DejaVu Sans Mono fonts, or provide the "
                       "complete Windows Arial/Consolas families. No Helvetica/Courier fallback "
                       "is used because it can silently lose mathematical characters.")


def styles():
    files, _, _ = _font_files()
    for name, path in zip(("BodyFont", "BodyFont-Bold", "BodyFont-Italic", "BodyFont-BoldItalic", "CodeFont"), files):
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily("BodyFont", normal="BodyFont", bold="BodyFont-Bold",
                                 italic="BodyFont-Italic", boldItalic="BodyFont-BoldItalic")
    face, bold, mono = "BodyFont", "BodyFont-Bold", "CodeFont"
    s = getSampleStyleSheet()
    for name in ["Normal", "BodyText"]:
        s[name].fontName = face; s[name].fontSize = 10; s[name].leading = 14
        s[name].spaceAfter = 7
    for name, size in [("Title", 27), ("Heading1", 20), ("Heading2", 15), ("Heading3", 12)]:
        s[name].fontName = bold; s[name].fontSize = size; s[name].leading = size * 1.25
        s[name].textColor = colors.HexColor("#" + NAVY); s[name].spaceBefore = 12; s[name].spaceAfter = 8
    s.add(ParagraphStyle("SmallNote", parent=s["BodyText"], fontSize=8, leading=11, textColor=colors.HexColor("#526574")))
    s.add(ParagraphStyle("CodeBlock", parent=s["BodyText"], fontName=mono, fontSize=7.5, leading=10,
                         backColor=colors.HexColor("#F1F4F6"), borderPadding=6))
    s.add(ParagraphStyle("ListText", parent=s["BodyText"], leftIndent=16,
                         firstLineIndent=0, bulletIndent=0))
    s.add(ParagraphStyle("TableText", parent=s["SmallNote"], fontSize=8, leading=10.5,
                         textColor=colors.HexColor("#" + INK)))
    s.add(ParagraphStyle("CompactTableText", parent=s["TableText"], fontSize=7.5, leading=9.5))
    return s


def _public_link(target, source=None):
    target = html.unescape(target)
    if target.startswith(("https://", "http://")):
        return target
    if source is None or target.startswith("#"):
        return None
    relative, _, fragment = target.partition("#")
    try:
        path = (source.parent / relative).resolve().relative_to(ROOT.resolve())
    except ValueError:
        return None
    return ("https://github.com/JoshGutierrez56/personalized-model-portfolios/blob/main/"
            + quote(path.as_posix()) + ("#" + quote(fragment) if fragment else ""))


def inline(s, source=None):
    s = html.escape(s)
    def link(match):
        target = _public_link(match[2], source)
        return f'<link href="{html.escape(target, quote=True)}" color="#{TEAL}">{match[1]}</link>' if target else match[1]
    s = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", link, s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", s)
    s = re.sub(r"`([^`]+)`", r'<font name="CodeFont">\1</font>', s)
    return s


def _word_inline(paragraph, text, source, mono):
    """Preserve bold, code and clickable links without raw Markdown in Word."""
    token = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))")
    for part in token.split(text):
        if not part:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\(([^)]+)\)", part)
        if link and (target := _public_link(link[2], source)):
            relationship = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
            hyperlink = OxmlElement("w:hyperlink"); hyperlink.set(qn("r:id"), relationship)
            run = OxmlElement("w:r"); props = OxmlElement("w:rPr")
            color = OxmlElement("w:color"); color.set(qn("w:val"), TEAL); props.append(color)
            underline = OxmlElement("w:u"); underline.set(qn("w:val"), "single"); props.append(underline)
            run.append(props); value = OxmlElement("w:t"); value.text = link[1]; run.append(value)
            hyperlink.append(run); paragraph._p.append(hyperlink)
        elif part.startswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            paragraph.add_run(part[1:-1]).font.name = mono
        else:
            paragraph.add_run(link[1] if link else part)


def _check_glyphs(text, font_name, context):
    supported = pdfmetrics.getFont(font_name).face.charWidths
    missing = sorted({character for character in text if not character.isspace()
                      and ord(character) not in supported})
    if missing:
        examples = ", ".join(f"{character} (U+{ord(character):04X})" for character in missing[:12])
        raise RuntimeError(f"{context}: {font_name} lacks required glyphs: {examples}. "
                           "Use a Unicode font with coverage rather than publishing missing-glyph boxes.")


def footer(canvas, doc):
    canvas.saveState(); w, h = doc.pagesize
    canvas.setStrokeColor(colors.HexColor("#D9E2E9")); canvas.line(45, 38, w - 45, 38)
    canvas.setFont("BodyFont", 8); canvas.setFillColor(colors.HexColor("#526574"))
    canvas.drawString(45, 26, "Josh Gutierrez | Portfolio Validation & Investment Oversight")
    canvas.drawRightString(w - 45, 26, str(doc.page)); canvas.restoreState()


class GuideDocTemplate(SimpleDocTemplate):
    """Expose section navigation in PDF readers without a second source of text."""
    def afterFlowable(self, flowable):
        if hasattr(flowable, "_outline_level"):
            self._bookmark_number = getattr(self, "_bookmark_number", 0) + 1
            key = f"section-{self._bookmark_number}"
            self.canv.bookmarkPage(key)
            self.canv.addOutlineEntry(flowable.getPlainText(), key,
                                      level=flowable._outline_level, closed=False)


def render_markdown(source: Path, stem: str):
    text = source.read_text(encoding="utf-8")
    if re.search(r"\bTODO\b|\bTBD\b|\[INSERT", text):
        raise ValueError(f"Unresolved drafting placeholder in {source.name}")
    st = styles(); flow = []; docx = Document()
    _, word_face, word_mono = _font_files()
    _check_glyphs(text, "BodyFont", source.name)
    sec = docx.sections[0]; sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65)
    sec.left_margin = Inches(.7); sec.right_margin = Inches(.7)
    normal = docx.styles["Normal"]; normal.font.name = word_face; normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for name in ["Title", "Heading 1", "Heading 2", "Heading 3"]:
        docx.styles[name].font.name = word_face
        docx.styles[name].font.color.rgb = RGBColor.from_string(NAVY)
    docx.core_properties.author = "Josh Gutierrez"
    docx.core_properties.title = stem.replace("_", " ")
    docx.core_properties.comments = ""
    heading = sec.header.paragraphs[0]
    heading.add_run("PORTFOLIO OVERSIGHT LAB").font.size = Pt(8)
    foot = sec.footer.paragraphs[0]
    foot.add_run("Josh Gutierrez | " + stem.replace("_", " ") + " | Page ").font.size = Pt(8)
    field = OxmlElement("w:fldSimple"); field.set(qn("w:instr"), "PAGE"); foot._p.append(field)
    code = []; in_code = False; table = []
    def flush_table():
        nonlocal table
        if not table: return
        n = max(map(len, table))
        table_style = st["CompactTableText"] if n >= 6 else st["TableText"]
        numeric_style = ParagraphStyle("NumericTable", parent=table_style, alignment=TA_RIGHT)
        def numeric(value):
            return re.fullmatch(r"[+\-−]?\$?\d[\d,]*(?:\.\d+)?%?", value.strip()) is not None
        numeric_columns = {
            i for i in range(n)
            if any(i < len(row) and numeric(row[i]) for row in table[1:])
            and all(i >= len(row) or numeric(row[i]) or row[i].strip() in {"", "—", "–", "-"}
                    for row in table[1:])
        }
        data = [[Paragraph(("<b>" if j == 0 else "") + inline(x.strip(), source)
                           + ("</b>" if j == 0 else ""), numeric_style if i in numeric_columns else table_style)
                 for i, x in enumerate(row)] for j, row in enumerate(table)]
        data = [r + [Paragraph("", table_style)] * (n - len(r)) for r in data]
        widths = [(7.05 * inch) / n] * n
        if n == 3 and any("Location" in cell for cell in table[0]):
            widths = [7.05 * inch * fraction for fraction in (.24, .49, .27)]
        elif n == 7:
            widths = [7.05 * inch * fraction for fraction in (.19, .135, .135, .135, .135, .135, .135)]
        elif n == 4 and table[0][0].strip() == "Minutes":
            widths = [7.05 * inch * fraction for fraction in (.12, .32, .10, .46)]
        t = Table(data, repeatRows=1, colWidths=widths)
        t.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#E8EFF3")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#D5DEE5")),
            ("LEFTPADDING", (0, 0), (-1, -1), 5), ("RIGHTPADDING", (0, 0), (-1, -1), 5)]))
        flow.extend([t, Spacer(1, 8)])
        dt = docx.add_table(rows=0, cols=n); dt.style = "Light Shading Accent 1"
        dt.autofit = False
        for i, width in enumerate(widths):
            dt.columns[i].width = Inches(width / inch)
        for row_number, row in enumerate(table):
            dr = dt.add_row(); cells = dr.cells
            if row_number == 0:
                repeat = OxmlElement("w:tblHeader"); dr._tr.get_or_add_trPr().append(repeat)
            for i, val in enumerate(row):
                cells[i].width = Inches(widths[i] / inch)
                p = cells[i].paragraphs[0]
                if i in numeric_columns:
                    p.alignment = 2  # WordprocessingML right alignment.
                _word_inline(p, val.strip(), source, word_mono)
                p.paragraph_format.space_after = Pt(3)
                for run in p.runs:
                    run.font.size = Pt(7.5 if n >= 6 else 8)
                    if row_number == 0: run.bold = True
        table = []
    for line in text.splitlines():
        if line.startswith("```"):
            flush_table()
            if in_code:
                raw = "\n".join(code)
                _check_glyphs(raw, "CodeFont", source.name + " code block")
                # Avoid off-page monospace: wrap only visually, preserve source separately.
                import textwrap
                wrapped = "\n".join(part for ln in code for part in (textwrap.wrap(ln, width=98, replace_whitespace=False) or [""]))
                flow.extend([Preformatted(wrapped, st["CodeBlock"]), Spacer(1, 8)])
                p = docx.add_paragraph(raw); p.style = "No Spacing"
                for run in p.runs: run.font.name = word_mono; run.font.size = Pt(8)
                code = []
            in_code = not in_code; continue
        if in_code: code.append(line); continue
        if line.startswith("|"):
            if not re.match(r"^\|[\s:|\-]+\|$", line): table.append(line.strip("|").split("|"))
            continue
        flush_table()
        if not line.strip(): continue
        if line.strip() == "---": flow.append(Spacer(1, 8)); continue
        m = re.match(r"^(#{1,6})\s+(.*)", line)
        if m:
            level, title = min(len(m[1]), 3), m[2]
            style = "Title" if len(m[1]) == 1 else f"Heading{level - 1}"
            heading = Paragraph(inline(title, source), st[style])
            heading._outline_level = max(0, min(len(m[1]) - 2, 1))
            flow.append(heading); docx.add_heading(title, level=0 if len(m[1]) == 1 else level - 1)
        else:
            bullet = re.match(r"^([-*]|\d+\.)\s+(.*)", line)
            content = bullet[2] if bullet else line.removeprefix("> ")
            for code_text in re.findall(r"`([^`]+)`", content):
                _check_glyphs(code_text, "CodeFont", source.name + " inline code")
            marker = (bullet[1] if bullet[1][0].isdigit() else "•") if bullet else None
            flow.append(Paragraph(inline(content, source), st["ListText"] if bullet else st["BodyText"], bulletText=marker))
            p = docx.add_paragraph(style="Normal")
            if marker:
                p.paragraph_format.left_indent = Pt(14)
                p.paragraph_format.first_line_indent = Pt(-14)
                p.add_run(marker + " ")
            _word_inline(p, content, source, word_mono)
    flush_table()
    pdf = OUT / f"{stem}.pdf"
    GuideDocTemplate(str(pdf), pagesize=(612, 792), rightMargin=45, leftMargin=45,
        topMargin=43, bottomMargin=50, title=stem.replace("_", " "), author="Josh Gutierrez").build(flow, onFirstPage=footer, onLaterPages=footer)
    docx.save(OUT / f"{stem}.docx")
    return {"pdf": pdf.name, "docx": f"{stem}.docx", "source_sha256": hashlib.sha256(source.read_bytes()).hexdigest()}


def memo():
    bundle = validate_release_bundle(ROOT)
    b, i, r = bundle["baseline"], bundle["incidents"], bundle["receipt"]
    st = styles()
    f = [Paragraph("Portfolio oversight", st["Title"]), Paragraph("Decision memo | independent validation and implementation economics", st["Heading2"]),
        Paragraph(f"Research evidence run: {r['run_id']} · {r['completed_utc'][:10]}", st["SmallNote"]),
        Paragraph("Recommendation", st["Heading2"]),
        Paragraph(f"Release the research evidence for review: baseline {b['status']}; {r['detected_cases']} of {r['injected_cases']} deliberate faults blocked; corrected inputs {r['corrected_retest_status']}. Do not interpret technical release as approval to trade a real account.", st["BodyText"]),
        Paragraph("The investment decision", st["Heading2"]),
        Paragraph("For the No gold account, lower-turnover allocation cut annual one-way turnover from 18.23% to 7.25%, but raised realized gross tracking error from 1.37% to 1.64%. At the assumed 5bp per dollar bought or sold, the annual cost-rate saving is 1.10bp of NAV: about $110 per constant $1m. That is a modest cost benefit for about 27bp more tracking risk. Prefer the minimum-risk implementation if tracking fidelity is the client's primary objective; accept lower turnover only when the mandate places a defensible value on avoiding trades. Taxes are not modeled.", st["BodyText"]),
        Paragraph("What was independently challenged", st["Heading2"]),
        Paragraph("A separate contract and challenger reconstruct total-return sleeve holdings, return contributions, self-financing costs and benchmark wealth from source returns and recorded trades. SQL checks data grain and trade/cost identities. The checker does not call the production optimizer's validation or accounting functions. Shared economic assumptions and shared source inputs remain a limitation of independence.", st["BodyText"]),
        Paragraph("Decision clock", st["Heading2"]),
        Paragraph("Estimate risk and propose weights at the previous close. Old holdings earn the execution day's return. Rebalance at that day's close and deduct costs from the account. Enforce weight bounds after execution; monitor, rather than prohibit, drift between rebalances. Realized TE is an outcome, not the same object as the ex-ante risk budget.", st["BodyText"]),
        Paragraph("Data and evidence", st["Heading2"]),
        Paragraph("Actual WRDS/CRSP CIZ data: eight ETFs and 32,192 source observations across 2010–2025; 2014–2025 evaluation, 30 account/method paths and 4,320 decisions. Current-vintage and retrospectively selected universe; the history has already been examined. All displayed accounts and trading-cost assumptions are hypothetical.", st["BodyText"]),
        PageBreak(), Paragraph("Failure investigation and release conditions", st["Heading1"])]
    for c in i["cases"]:
        f.append(Paragraph(f"{c['id']} | {html.escape(c['title'])} — {c['observed_status']}", st["Heading3"]))
        f.append(Paragraph(inline(c["injection"] + " " + c["corrective_action"] + " Corrected-input retest: " + c["corrected_retest"] + "."), st["BodyText"]))
    f.extend([Paragraph("What the evidence does not establish", st["Heading2"]),
        Paragraph("Faults are deliberately injected copies, not historical production incidents. The wrong-return case holds recorded trades fixed; its replay is diagnostic inconsistency, not a self-consistent alternative strategy. The experiment does not establish live execution readiness, tax-lot handling, stock-level factor attribution, market-impact calibration, multi-currency accounting or future alpha. Contribution reconciliation explains the accounting of this model, not a causal factor model. Passing these checks does not certify all possible errors.", st["BodyText"]),
        Paragraph("Next release conditions", st["Heading2"]),
        Paragraph("For production use: reconcile to an independent custodian/order source; implement corporate-action and tax-lot policies, market-impact and execution controls, operational ownership, monitoring and deployment review. For research discussion: retain run hashes, tolerance definitions, exception acceptance criteria and the corrected retest receipt. Detailed licensed security-day data remain private.", st["BodyText"]),
        Paragraph("Evidence: validation_summary.json, incident_summary.json, run_receipt.json, historical_stress.json and Portfolio_Oversight.xlsx. Created with AI coding and drafting assistance; interpretation and technical explanations require the presenter's own understanding.", st["SmallNote"])])
    path = OUT / "Portfolio_Oversight_Memo.pdf"
    SimpleDocTemplate(str(path), pagesize=(612, 792), leftMargin=45, rightMargin=45, topMargin=38, bottomMargin=48,
                      title="Portfolio Oversight Decision Memo", author="Josh Gutierrez").build(f, onFirstPage=footer, onLaterPages=footer)
    return path.name


def main():
    validate_release_bundle(ROOT)
    OUT.mkdir(parents=True, exist_ok=True)
    evidence = {"created_utc": datetime.now(timezone.utc).isoformat(), "workbook": workbook(), "memo": memo()}
    for path, stem in [(ROOT / "docs/INTERVIEW_GUIDE.md", "Interview_Guide"),
                       (ROOT / "docs/OVERSIGHT_METHODOLOGY.md", "Oversight_Methodology"),
                       (ROOT / "docs/INTERVIEW_PRACTICE.md", "Interview_Practice")]:
        if path.exists(): evidence[stem] = render_markdown(path, stem)
    _, body_font, code_font = _font_files()
    evidence["fonts"] = {"pdf_body": body_font, "pdf_code": code_font,
                         "pdf_fonts_embedded": True,
                         "docx_note": "Word files request these families but do not embed fonts; the reader may substitute available fonts."}
    evidence["files"] = {p.name: {"bytes": p.stat().st_size, "sha256": hashlib.sha256(p.read_bytes()).hexdigest()}
        for p in sorted(OUT.iterdir()) if p.is_file() and p.suffix in {".pdf", ".docx", ".xlsx"}}
    (OUT / "report_manifest.json").write_text(json.dumps(evidence, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(evidence, indent=2))


if __name__ == "__main__":
    main()
